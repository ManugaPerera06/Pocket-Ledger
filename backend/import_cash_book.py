"""Import the supplied 2023–2026 cash book into the normalized SQLite table."""
from __future__ import annotations

import argparse
import re
import sys
from contextlib import closing
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import column_index_from_string

# Support running this file directly from the repository root or as a module.
ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
from backend.db import get_connection, init_db  # noqa: E402

SHEET_CONFIG = {
    "2023": {"date": "A", "details": "B", "amount": "H", "min_row": 4},
    "2024": {"date": "A", "details": "B", "income": "H", "expense": "J", "min_row": 4},
    "2025": {"date": "A", "details": "B", "income": "H", "expense": "J", "min_row": 4},
    "2026": {"date": "A", "details": "B", "income": "C", "expense": "D", "min_row": 3},
}
SKIP_PATTERNS = re.compile(r"^\s*(all incomes?|all expenditures?|year profit|details)\s*$", re.I)
CATEGORY_RULES = [
    ("Savings", ["purse"]),
    ("Education & Classes", ["clz", "class", "iit", "sakya", " ism", "ucl", "ssb", "ccc", "maths", "ict", "econ", "commerce", "accounting", "driving", "learners", "school", "career fair", "swimming", "gym"]),
    ("Food", ["kottu", "roti", "parata", "bun", "bread", "egg", "chocolate", "ice cream", "omlet", "soup", "coconut", "chicken puff", "dinner", "lunch"]),
    ("Shopping", ["clothes", "slippers", "t-shirt", "shirt", "book", "gift", "data card", "figma", "tailors"]),
    ("Health", ["doctor", "hospital", "medical", "pampus"]),
    ("Outings & Social", ["party", "mall", "havelock", "bmich", "match", "cemetery", "nugegoda", "dehiwala", "maharagama", "'s place", "dialog", "photo", "cool planet", "awurudu", "b'day", "uthsavaya", "semina"]),
    ("Family", ["amma", "thaththa", "achchi", "punchi", "seeya", "mahappa", "malli"]),
    ("Loans", ["loan"]),
    ("Personal Care", ["hair"]),
]


def category_for(description: str) -> str:
    text = description.lower()
    for category, keywords in CATEGORY_RULES:
        if any(keyword in text for keyword in keywords):
            return category
    return "Uncategorized"


def parse_date(raw: Any) -> str:
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    return datetime.strptime(str(raw).strip(), "%d.%m.%y").date().isoformat()


def amount_cents(value: Any) -> int | None:
    if not isinstance(value, (int, float)) or value == 0:
        return None
    return int((Decimal(str(abs(value))) * 100).quantize(Decimal("1")))


def extract_rows(ws: Any, config: dict[str, Any], sheet_year: str, source_file: str) -> list[dict[str, Any]]:
    indexes = {key: column_index_from_string(letter) - 1 for key, letter in config.items() if key != "min_row"}
    rows: list[dict[str, Any]] = []
    last_date: str | None = None

    def value(row: tuple[Any, ...], key: str) -> Any:
        index = indexes[key]
        return row[index].value if index < len(row) else None

    for row in ws.iter_rows(min_row=config["min_row"]):
        raw_description = value(row, "details")
        if not raw_description or not isinstance(raw_description, str) or SKIP_PATTERNS.search(raw_description):
            continue
        raw_date = value(row, "date")
        if raw_date:
            try:
                last_date = parse_date(raw_date)
            except ValueError:
                continue
            if not last_date.startswith(sheet_year):
                raise ValueError(f"{sheet_year}!{row[0].row}: date {last_date} is outside the sheet year")
        if not last_date:
            continue

        if "amount" in config:
            raw_amount = value(row, "amount")
            cents = amount_cents(raw_amount)
            tx_type = "INCOME" if isinstance(raw_amount, (int, float)) and raw_amount > 0 else "EXPENSE"
        else:
            income = amount_cents(value(row, "income"))
            expense = amount_cents(value(row, "expense"))
            if income:
                cents, tx_type = income, "INCOME"
            elif expense:
                cents, tx_type = expense, "EXPENSE"
            else:
                continue
        if not cents:
            continue
        rows.append({
            "transaction_date": last_date,
            "description": raw_description.strip(),
            "category": category_for(raw_description),
            "type": tx_type,
            "amount_cents": cents,
            "source_file": source_file,
            "source_sheet": sheet_year,
            "source_row": row[0].row,
            "source_raw_date": str(raw_date).strip() if raw_date else None,
        })
    return rows


def import_workbook(path: str | Path, replace: bool = False) -> dict[str, int]:
    workbook_path = Path(path)
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    all_rows: list[dict[str, Any]] = []
    counts: dict[str, int] = {}
    for sheet_name, config in SHEET_CONFIG.items():
        if sheet_name not in workbook.sheetnames:
            continue
        parsed = extract_rows(workbook[sheet_name], config, sheet_name, workbook_path.name)
        counts[sheet_name] = len(parsed)
        all_rows.extend(parsed)
    init_db()
    with closing(get_connection()) as conn:
        if replace:
            conn.execute("DELETE FROM transactions")
        conn.executemany(
            """INSERT INTO transactions
            (transaction_date, description, category, type, amount_cents, source_file, source_sheet, source_row, source_raw_date)
            VALUES (:transaction_date, :description, :category, :type, :amount_cents, :source_file, :source_sheet, :source_row, :source_raw_date)""",
            all_rows,
        )
        conn.commit()
    counts["total"] = len(all_rows)
    return counts


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx_path", type=Path)
    parser.add_argument("--replace", action="store_true", help="Clear existing transactions before import")
    args = parser.parse_args()
    print(import_workbook(args.xlsx_path, replace=args.replace))


if __name__ == "__main__":
    main()
