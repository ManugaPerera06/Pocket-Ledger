"""
One-off migration script: imports historical pocket-money data from
My Cash Book.xlsx (sheets 2023-2026) into the transactions table.
"""
import argparse
import re
from contextlib import closing
from datetime import datetime

import openpyxl
from openpyxl.utils import column_index_from_string

from db import get_connection, init_db

# Each sheet has a slightly different layout, so each gets its own column map.
SHEET_CONFIG = {
    "2023": {"date": "A", "details": "B", "amount": "H", "min_row": 4},
    "2024": {"date": "A", "details": "B", "income": "H", "expense": "J", "min_row": 4},
    "2025": {"date": "A", "details": "B", "income": "H", "expense": "J", "min_row": 4},
    "2026": {"date": "A", "details": "B", "income": "C", "expense": "D", "min_row": 3},
}

# Rows that are weekly/monthly summaries, not real transactions, get skipped.
SKIP_PATTERNS = re.compile(
    r"balance|profit|loss|all income|all expenditure|year profit|^details$",
    re.IGNORECASE,
)

# First matching rule wins; anything unmatched becomes "Uncategorized".
CATEGORY_RULES = [
    ("Savings", ["purse"]),
    ("Education & Classes", ["clz", "class", "iit", "sakya", " ism", "ucl", "ssb", "ccc",
                              "maths", "ict", "econ", "commerce", "accounting", "driving",
                              "learners", "school", "career fair", "swimming", "gym"]),
    ("Food", ["kottu", "roti", "parata", "bun", "bread", "egg", "chocolate",
              "ice cream", "omlet", "soup", "coconut", "chicken puff", "dinner", "lunch"]),
    ("Shopping", ["clothes", "slippers", "t-shirt", "shirt", "book", "gift",
                  "data card", "figma", "tailors"]),
    ("Health", ["doctor", "hospital", "medical", "pampus"]),
    ("Outings & Social", ["party", "mall", "havelock", "bmich", "match", "cemetery",
                           "nugegoda", "dehiwala", "maharagama", "'s place", "dialog",
                           "photo", "cool planet", "awurudu", "b'day", "uthsavaya", "semina"]),
    ("Family", ["amma", "thaththa", "achchi", "punchi", "seeya", "mahappa", "malli"]),
    ("Loans", ["loan"]),
    ("Personal Care", ["hair"]),
]


def category_for(details: str) -> str:
    text = details.lower()
    for category, keywords in CATEGORY_RULES:
        if any(kw in text for kw in keywords):
            return category
    return "Uncategorized"


def parse_date(raw) -> str:
    """Accepts 'DD.MM.YY' text or a real datetime cell; returns 'YYYY-MM-DD'."""
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    return datetime.strptime(str(raw).strip(), "%d.%m.%y").date().isoformat()


def extract_rows(ws, config):
    rows = []
    last_date = None

    # Resolve column letters to 0-based positions once, up front.
    col_idx = {
        key: column_index_from_string(letter) - 1
        for key, letter in config.items() if key != "min_row"
    }

    def val(row, key):
        idx = col_idx[key]
        return row[idx].value if idx < len(row) else None

    for row in ws.iter_rows(min_row=config["min_row"]):
        details = val(row, "details")
        if not details or not isinstance(details, str) or SKIP_PATTERNS.search(details):
            continue

        date_val = val(row, "date")
        if date_val:
            try:
                last_date = parse_date(date_val)
            except ValueError:
                continue
        if last_date is None:
            continue  # stray row before any date has been seen

        if "amount" in config:  # 2023 style: single signed column
            amt = val(row, "amount")
            if not isinstance(amt, (int, float)) or amt == 0:
                continue
            tx_type = "INCOME" if amt > 0 else "EXPENSE"
            cents = abs(round(amt * 100))
        else:  # income/expense split, expense stored as negative
            income = val(row, "income")
            expense = val(row, "expense")
            if isinstance(income, (int, float)) and income != 0:
                tx_type, cents = "INCOME", abs(round(income * 100))
            elif isinstance(expense, (int, float)) and expense != 0:
                tx_type, cents = "EXPENSE", abs(round(expense * 100))
            else:
                continue

        rows.append({
            "type": tx_type,
            "amount_cents": int(cents),
            "category": category_for(details),
            "note": details.strip(),
            "transaction_date": last_date,
        })
    return rows


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx_path")
    parser.add_argument("--dry-run", action="store_true",
                         help="Preview counts/sample rows without writing to the database.")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.xlsx_path, data_only=True)

    all_rows = []
    for sheet_name, config in SHEET_CONFIG.items():
        if sheet_name not in wb.sheetnames:
            print(f"  (sheet '{sheet_name}' not found, skipping)")
            continue
        sheet_rows = extract_rows(wb[sheet_name], config)
        total = sum(r["amount_cents"] for r in sheet_rows) / 100
        print(f"{sheet_name}: {len(sheet_rows)} transactions parsed "
              f"({sum(1 for r in sheet_rows if r['type']=='INCOME')} income, "
              f"{sum(1 for r in sheet_rows if r['type']=='EXPENSE')} expense)")
        all_rows.extend(sheet_rows)

    all_rows.sort(key=lambda r: r["transaction_date"])
    print(f"\nTotal: {len(all_rows)} transactions")

    if args.dry_run:
        print("\nSample rows:")
        for r in all_rows[:10]:
            print(" ", r)
        print("\n(dry run — nothing written to the database)")
        return

    init_db()
    with closing(get_connection()) as conn, conn:
        conn.executemany(
            """
            INSERT INTO transactions (type, amount_cents, category, note, transaction_date)
            VALUES (:type, :amount_cents, :category, :note, :transaction_date)
            """,
            all_rows,
        )
    print("Imported successfully.")


if __name__ == "__main__":
    main()